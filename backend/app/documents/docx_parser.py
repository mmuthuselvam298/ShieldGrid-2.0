import logging
from pathlib import Path
from typing import Dict, Any, List
import docx

logger = logging.getLogger(__name__)

class DocxParser:
    @staticmethod
    def parse(file_path: Path) -> Dict[str, Any]:
        """
        Extract text from DOCX paragraphs and tables.
        """
        doc = docx.Document(str(file_path))
        lines = []

        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(p.text)

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    lines.append(" | ".join(row_text))

        full_text = "\n".join(lines)
        return {
            "page_count": 1, # DOCX does not have fixed pages without layout engine
            "text": full_text,
            "is_scanned": False
        }

    @staticmethod
    def redact(input_path: Path, output_path: Path, findings_to_redact: List[Dict[str, Any]]) -> int:
        doc = docx.Document(str(input_path))
        applied_count = 0

        # Replace text in paragraphs
        for item in findings_to_redact:
            target = item.get("text", "")
            replacement = item.get("redacted_value", "████████")
            if not target:
                continue

            for p in doc.paragraphs:
                if target in p.text:
                    p.text = p.text.replace(target, replacement)
                    applied_count += 1

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if target in cell.text:
                            cell.text = cell.text.replace(target, replacement)
                            applied_count += 1

        doc.save(str(output_path))
        return applied_count

docx_parser = DocxParser()
